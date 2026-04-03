import openpyxl
import os
from datetime import datetime

EXCEL_FILE = "לקוחות_רועי.xlsx"
PROPERTIES_FILE = "properties.xlsx"

# עמודות לאנליטיקה פנימית: Date, User, Type, Neighborhood, Address, Price, Sqm, Floor, AI Score, Publisher Type, URL
PROPERTY_HEADERS = [
    "תאריך",
    "משתמש",
    "סוג (השכרה/מכירה)",
    "שכונה",
    "כתובת",
    "מחיר",
    'מ"ר',
    "קומה",
    "ציון AI",
    "סוג מפרסם (פרטי/תיווך)",
    "קישור",
]


def init_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "לקוחות"
        headers = ["תאריך", "שם", "אזור", "חדרים", "תקציב", "סוג", "סטטוס", "הערות"]
        ws.append(headers)
        for col in range(1, 9):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18
        wb.save(EXCEL_FILE)
        print(f"קובץ אקסל נוצר: {EXCEL_FILE}")


def init_properties_excel():
    if not os.path.exists(PROPERTIES_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "נכסים"
        ws.append(PROPERTY_HEADERS)
        for col in range(1, len(PROPERTY_HEADERS) + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 22
        wb.save(PROPERTIES_FILE)


def add_client(name: str, area: str, rooms: str, budget: str, deal_type: str = "לא ידוע", notes: str = ""):
    init_excel()
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb.active

    row = [
        datetime.now().strftime("%d/%m/%Y %H:%M"),
        name,
        area,
        rooms,
        budget,
        deal_type,
        "חדש",
        notes
    ]

    ws.append(row)
    wb.save(EXCEL_FILE)
    print(f"לקוח נוסף לאקסל: {name}")


def save_to_excel(data, user_name: str):
    """שומר נתוני נכסים ל-properties.xlsx — לניתוח מגמות ומחירים.

    data: dict בודד או רשימת dict עם מפתחות אופציונליים:
    date, type, neighborhood, address, price, sqm, floor, ai_score, publisher_type, url
    """
    init_properties_excel()
    wb = openpyxl.load_workbook(PROPERTIES_FILE)
    ws = wb.active
    now = datetime.now().strftime("%d/%m/%Y %H:%M")

    rows = data if isinstance(data, list) else [data]
    for p in rows:
        ws.append([
            p.get("date") or now,
            user_name,
            p.get("type", ""),
            p.get("neighborhood", ""),
            p.get("address", ""),
            p.get("price", ""),
            p.get("sqm", ""),
            p.get("floor", ""),
            p.get("ai_score", ""),
            p.get("publisher_type", ""),
            p.get("url", ""),
        ])
    wb.save(PROPERTIES_FILE)
    print(f"נשמרו {len(rows)} נכסים ל-{PROPERTIES_FILE} עבור {user_name}")
